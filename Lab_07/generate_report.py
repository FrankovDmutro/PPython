"""Script to generate the official DOCX laboratory report for Lab 07, Variant 15."""

from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def create_report():
    doc = Document()

    # Base page settings
    for sec in doc.sections:
        sec.top_margin = Inches(0.79)
        sec.bottom_margin = Inches(0.79)
        sec.left_margin = Inches(0.98)
        sec.right_margin = Inches(0.79)

    # Configure Normal style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    def add_p(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, first_line=0.49):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.first_line_indent = Inches(first_line) if first_line > 0 else 0
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0, 32, 96)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(31, 78, 121)
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F2F2F2")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = 0
        run = p.add_run(code_text.strip())
        run.font.name = 'Courier New'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(20, 20, 20)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== ТИТУЛЬНА СТОРІНКА ====================
    add_p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    add_p("Кафедра інформаційних технологій та програмування", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)

    for _ in range(3):
        doc.add_paragraph()

    add_p("ЗВІТ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    add_p("з лабораторної роботи № 7", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    add_p("з дисципліни «Програмування на Python»", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    add_p("на тему: «Робота з базами даних у Python»", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)

    for _ in range(4):
        doc.add_paragraph()

    add_p("Індивідуальний варіант № 15", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)
    add_p("Виконав: студент Франков", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)

    for _ in range(5):
        doc.add_paragraph()

    add_p("2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=0)
    doc.add_page_break()

    # ==================== ТЕМА ТА МЕТА РОБОТИ ====================
    add_h1("1. ТЕМА ТА МЕТА РОБОТИ")
    add_p("Тема роботи: Робота з базами даних у Python.")
    add_p("Мета роботи: Ознайомитися з основними принципами побудови та організації шару зберігання даних (Persistence Layer) у Python-застосунках. Навчитися працювати з реляційними базами даних, мовою структурованих запитів SQL, вбудованою СУБД SQLite, офіційним стандартом Python DB-API (sqlite3), безпечними параметризованими запитами, реалізацією операцій CRUD (Create, Read, Update, Delete), керуванням транзакціями та принципами ACID (Atomicity, Consistency, Isolation, Durability). Опанувати патерн проєктування Repository для ізоляції бізнес-логіки від деталей сховища, об'єктно-реляційне відображення SQLAlchemy 2.0 ORM, роботу зі зв'язками між сутностями (One-to-Many), міграції схеми бази даних за допомогою інструменту Alembic, а також написання автоматизованих інтеграційних тестів із використанням тестової бази даних в оперативній пам'яті.")

    # ==================== ПОСТАНОВКА ЗАДАЧІ ====================
    add_h1("2. ПОСТАНОВКА ЗАДАЧІ (ВАРІАНТ № 15)")
    add_p("Згідно з індивідуальним варіантом № 15 предметною областю є: «Облік результатів наукових експериментів» (розширення та інтеграція результатів попередніх робіт).")
    add_p("Вимоги варіанта № 15:")
    add_p("• Продовжити проєкт попередніх робіт. Для кожного експерименту зберігати обов'язкові атрибути: назву (title), дослідника (researcher), вимірюваний параметр (parameter), числовий результат (result) та дату проведення (date);")
    add_p("• Спроєктувати та створити реляційні таблиці researchers та experiments із коректними первинними (PRIMARY KEY) та зовнішніми ключами (FOREIGN KEY);")
    add_p("• Реалізувати шар доступу до даних на базі патерну Repository: класи ResearcherRepository та ExperimentRepository;")
    add_p("• Зберегти та розширити функціонал: додавання експерименту, пошук за дослідником, обчислення середнього, мінімального та максимального результатів, фільтрація за числовим діапазоном, повний цикл CRUD та сортування записів за датою (ORDER BY date);")
    add_p("• Засобами стандарту Python DB-API (модуль sqlite3) реалізувати параметризований пошук за дослідником та діапазоном значень (захист від SQL Injection), а також обчислення агрегатів AVG(), MIN(), MAX(), COUNT();")
    add_p("• Реалізувати транзакційні операції: пакетне додавання серії експериментів (Batch Insert) із гарантованим відкатом (rollback) у разі виникнення помилки або неприпустимого значення (наприклад, температура нижче абсолютного нуля);")
    add_p("• Створити та застосувати міграцію бази даних за допомогою Alembic для додавання стовпця notes (додаткові примітки);")
    add_p("• Розробити вичерпний набір інтеграційних тестів (Integration Tests) із перевіркою CRUD, зв'язків між таблицями, агрегатів, фільтрації, сортування та транзакційного відкату.")

    # ==================== ТЕОРЕТИЧНІ ВІДОМОСТІ ====================
    add_h1("3. ТЕОРЕТИЧНІ ВІДОМОСТІ")
    add_p("3.1. Реляційні бази даних та мова SQL. Реляційна база даних організовує дані у вигляді взаємопов'язаних таблиць, що складаються з рядків (записів) та стовпчиків (полів). Мова SQL (Structured Query Language) є декларативною мовою визначення структури даних (DDL: CREATE TABLE, ALTER TABLE) та маніпулювання даними (DML: SELECT, INSERT, UPDATE, DELETE). Первинний ключ (Primary Key) забезпечує унікальну ідентифікацію кожного рядка, а зовнішній ключ (Foreign Key) гарантує реляційну цілісність, встановлюючи зв'язок між рядками дочірньої та батьківської таблиць.")
    add_p("3.2. Стандарт Python DB-API та безпека запитів. Python Database API Specification (PEP 249) визначає уніфікований інтерфейс для доступу до СУБД через об'єкти Connection (з'єднання) та Cursor (курсор). Критичним правилом безпеки є використання параметризованих запитів (наприклад, cursor.execute('SELECT * FROM t WHERE name = ?', (val,))). Пряма конкатенація або використання f-strings створює вразливість до SQL Injection (ін'єкції зловмисного SQL-коду), що може призвести до несанкціонованого доступу або повного знищення бази даних.")
    add_p("3.3. Транзакції та принципи ACID. Транзакція — це логічна одиниця роботи, що об'єднує послідовність SQL-операцій. Транзакції гарантують надійність відповідно до вимог ACID:")
    add_p("• Atomicity (Атомарність): неподільність транзакції — виконуються або всі операції (commit), або жодна (rollback);")
    add_p("• Consistency (Узгодженість): перехід бази даних лише з одного валідного стану в інший зі збереженням усіх обмежень цілісності;")
    add_p("• Isolation (Ізольованість): паралельні транзакції не впливають одна на одну;")
    add_p("• Durability (Довговічність): зафіксовані зміни (commit) гарантовано зберігаються навіть у разі раптового збою живлення.")
    add_p("3.4. Патерн Repository та ORM SQLAlchemy. Repository Pattern інкапсулює логіку доступу до сховища, надаючи сервісному шару об'єктно-орієнтований інтерфейс, подібний до колекції у пам'яті. Це повністю відокремлює бізнес-правила від конкретної СУБД та SQL-діалекту. SQLAlchemy 2.0 є провідною бібліотекою Object-Relational Mapping (ORM) для Python, що зіставляє класи Python із таблицями реляційної бази даних, автоматизуючи генерацію запитів та керування зв'язками між сутностями (relationship).")
    add_p("3.5. Міграції бази даних (Alembic). Міграція — це версіонований скрипт, що описує перехід структури бази даних від однієї ревізії схеми до іншої (upgrade) та можливість повернення назад (downgrade). Інструмент Alembic дозволяє командно керувати еволюцією схеми бази даних без ручного виконання DDL-запитів.")

    # ==================== СТРУКТУРА ТА СХЕМА БАЗИ ДАНИХ ====================
    add_h1("4. АРХІТЕКТУРА ПРОЄКТУ ТА СХЕМА БАЗИ ДАНИХ")

    add_h2("4.1. Схема бази даних (Варіант № 15)")
    add_p("Спроєктовано дві реляційні таблиці, пов'язані відношенням One-to-Many (один дослідник може здійснювати багато наукових експериментів):")

    schema_table = doc.add_table(rows=8, cols=5)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Таблиця", "Стовпець", "Тип даних", "Ключ / Обмеження", "Призначення"]
    for i, h in enumerate(s_headers):
        cell = schema_table.cell(0, i)
        set_cell_background(cell, "D9E1F2")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)

    schema_data = [
        ("researchers", "id", "INTEGER", "PRIMARY KEY AUTOINCREMENT", "Унікальний ідентифікатор дослідника"),
        ("researchers", "name", "VARCHAR(100)", "NOT NULL, UNIQUE", "ПІБ дослідника (індексоване)"),
        ("researchers", "department", "VARCHAR(100)", "NOT NULL", "Науковий підрозділ / кафедра"),
        ("experiments", "id", "INTEGER", "PRIMARY KEY AUTOINCREMENT", "Унікальний ідентифікатор експерименту"),
        ("experiments", "title", "VARCHAR(200)", "NOT NULL", "Назва наукового експерименту"),
        ("experiments", "researcher_id", "INTEGER", "FOREIGN KEY -> researchers.id", "Зовнішній ключ (CASCADE DELETE)"),
        ("experiments", "parameter", "VARCHAR(50)", "NOT NULL", "Фізичний параметр (температура, тиск...)"),
    ]
    for row_idx, r in enumerate(schema_data, start=1):
        for col_idx, val in enumerate(r):
            cell = schema_table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in {0, 2, 3} else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_p("Таблиця experiments додатково містить стовпці:")
    add_p("• result: FLOAT NOT NULL — числове значення вимірювання;")
    add_p("• date: VARCHAR(20) NOT NULL — дата проведення експерименту в форматі ISO YYYY-MM-DD;")
    add_p("• notes: TEXT NULL — додаткові примітки та опис умов (додано через Alembic-міграцію).")

    add_h2("4.2. DDL-запити створення таблиць")
    ddl_text = """CREATE TABLE researchers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name VARCHAR(100) NOT NULL UNIQUE,
    department VARCHAR(100) NOT NULL
);

CREATE TABLE experiments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title VARCHAR(200) NOT NULL,
    researcher_id INTEGER NOT NULL REFERENCES researchers(id) ON DELETE CASCADE,
    parameter VARCHAR(50) NOT NULL,
    result FLOAT NOT NULL,
    date VARCHAR(20) NOT NULL,
    notes TEXT
);

CREATE INDEX ix_experiments_researcher_id ON experiments(researcher_id);
CREATE INDEX ix_experiments_parameter ON experiments(parameter);
CREATE INDEX ix_experiments_date ON experiments(date);"""
    add_code_block(ddl_text)

    # ==================== ОПИС РЕАЛІЗАЦІЇ ТА ЛІСТИНГИ ====================
    add_h1("5. ОПИС РЕАЛІЗАЦІЇ ТА ОСНОВНІ ЛІСТИНГИ КОДУ")

    add_h2("5.1. Налаштування бази даних (database.py)")
    add_p("Модуль database.py інкапсулює створення об'єктів Engine, фабрики сесій sessionmaker, базового класу декларативних моделей Base та контекстного менеджера безпечної транзакційної сесії get_session:")
    db_text = """class Base(DeclarativeBase):
    \"\"\"Declarative base class for all SQLAlchemy ORM models.\"\"\"

def get_engine(db_url: str | None = None, echo: bool = False) -> Engine:
    url = db_url or DEFAULT_DB_URL
    return create_engine(url, echo=echo, connect_args={"check_same_thread": False})

def get_session_factory(engine: Engine) -> sessionmaker[Session]:
    return sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

@contextmanager
def get_session(session_factory: sessionmaker[Session]) -> Generator[Session, None, None]:
    session = session_factory()
    try:
        yield session
        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        session.close()"""
    add_code_block(db_text)

    add_h2("5.2. ORM-моделі сутностей (models.py)")
    add_p("Використано сучасний синтаксис SQLAlchemy 2.0 (Mapped та mapped_column) зі встановленням зв'язку One-to-Many:")
    models_text = """class Researcher(Base):
    __tablename__ = "researchers"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    name: Mapped[str] = mapped_column(String(100), nullable=False, unique=True, index=True)
    department: Mapped[str] = mapped_column(String(100), nullable=False)

    experiments: Mapped[list["Experiment"]] = relationship(
        "Experiment", back_populates="researcher", cascade="all, delete-orphan", order_by="Experiment.id"
    )

class Experiment(Base):
    __tablename__ = "experiments"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    title: Mapped[str] = mapped_column(String(200), nullable=False)
    researcher_id: Mapped[int] = mapped_column(Integer, ForeignKey("researchers.id", ondelete="CASCADE"), nullable=False, index=True)
    parameter: Mapped[str] = mapped_column(String(50), nullable=False, index=True)
    result: Mapped[float] = mapped_column(Float, nullable=False)
    date: Mapped[str] = mapped_column(String(20), nullable=False, index=True)
    notes: Mapped[str | None] = mapped_column(Text, nullable=True)

    researcher: Mapped["Researcher"] = relationship("Researcher", back_populates="experiments")"""
    add_code_block(models_text)

    add_h2("5.3. Реалізація патерну Repository (repositories.py)")
    add_p("Репозиторії надають методи CRUD, фільтрації та розрахунку SQL-агрегатів через об'єкт Session:")
    repo_text = """class ExperimentRepository:
    def __init__(self, session: Session) -> None:
        self.session = session

    def create(self, title: str, researcher_id: int, parameter: str, result: float, date: str, notes: str | None = None) -> Experiment:
        experiment = Experiment(title=title.strip(), researcher_id=researcher_id, parameter=parameter.strip().lower(), result=float(result), date=date.strip(), notes=notes)
        self.session.add(experiment)
        self.session.flush()
        return experiment

    def find_by_range(self, min_val: float, max_val: float, parameter: str | None = None) -> list[Experiment]:
        stmt = select(Experiment).where(Experiment.result >= min_val, Experiment.result <= max_val)
        if parameter is not None:
            stmt = stmt.where(Experiment.parameter == parameter.strip().lower())
        stmt = stmt.order_by(Experiment.result.asc())
        return list(self.session.scalars(stmt).all())

    def get_aggregates(self, parameter: str | None = None) -> dict[str, float]:
        stmt = select(
            func.count(Experiment.id).label("count"),
            func.coalesce(func.avg(Experiment.result), 0.0).label("avg"),
            func.coalesce(func.min(Experiment.result), 0.0).label("min"),
            func.coalesce(func.max(Experiment.result), 0.0).label("max"),
        )
        if parameter is not None:
            stmt = stmt.where(Experiment.parameter == parameter.strip().lower())
        row = self.session.execute(stmt).one()
        return {"count": float(row.count), "average": float(row.avg), "min": float(row.min), "max": float(row.max)}"""
    add_code_block(repo_text)

    add_h2("5.4. Безпечні параметризовані запити DB-API (db_api.py)")
    add_p("Реалізація прямих запитів до SQLite через sqlite3 без використання f-string конкатенацій:")
    dbapi_text = """def db_api_search_by_researcher(conn: sqlite3.Connection, researcher_name: str) -> list[dict[str, Any]]:
    query = \"\"\"
        SELECT e.id, e.title, r.name AS researcher, r.department, e.parameter, e.result, e.date, e.notes
        FROM experiments e
        JOIN researchers r ON e.researcher_id = r.id
        WHERE r.name = ?
        ORDER BY e.date ASC;
    \"\"\"
    cursor = conn.cursor()
    cursor.execute(query, (researcher_name.strip(),))
    return [dict(row) for row in cursor.fetchall()]"""
    add_code_block(dbapi_text)

    add_h2("5.5. Транзакційний сервісний шар (services.py)")
    add_p("Метод batch_add_experiments_transactional гарантує атомарний відкат (rollback) всієї пачки експериментів, якщо хоча б один результат порушує фізичні межі:")
    serv_text = """def batch_add_experiments_transactional(self, records: list[dict[str, Any]]) -> list[Experiment]:
    created_experiments: list[Experiment] = []
    try:
        for item in records:
            res_val = float(item["result"])
            param_name = str(item["parameter"]).strip().lower()
            if param_name == "temperature" and res_val < -273.15:
                raise ValueError(f"Physical constraint violation: Temperature {res_val} < -273.15 C")

            researcher = self.register_researcher(item["researcher_name"], item.get("department", "General Physics"))
            exp = self.experiment_repo.create(item["title"], researcher.id, param_name, res_val, item["date"], item.get("notes"))
            created_experiments.append(exp)

        self.session.flush()
        return created_experiments
    except Exception:
        self.session.rollback()
        raise"""
    add_code_block(serv_text)

    add_h2("5.6. Міграція Alembic (001_add_notes_to_experiments.py)")
    add_p("Скрипт міграції використовує batch_alter_table для безпечної модифікації таблиць у СУБД SQLite:")
    mig_text = """def upgrade() -> None:
    with op.batch_alter_table("experiments", schema=None) as batch_op:
        batch_op.add_column(sa.Column("notes", sa.Text(), nullable=True))

def downgrade() -> None:
    with op.batch_alter_table("experiments", schema=None) as batch_op:
        batch_op.drop_column("notes")"""
    add_code_block(mig_text)

    # ==================== РЕЗУЛЬТАТИ ТЕСТУВАННЯ ТА РОБОТИ ====================
    add_h1("6. РЕЗУЛЬТАТИ РОБОТИ ТА ТЕСТУВАННЯ")

    add_h2("6.1. Демонстраційний запуск CLI застосунку")
    add_p("Виконано реальний запуск програми experiment_manager.main із підключенням до бази даних SQLite:")
    run_log = """Connecting to SQLite database: sqlite:///Lab_07/data/experiments.db
Database tables ('researchers', 'experiments') initialized successfully.

=== Experiments sorted by date ===
[2026-04-01] ID=1 | 'Cryogenic Calibration' | temperature: 24.5 | Researcher: Dr. Frankov
[2026-04-02] ID=2 | 'Vacuum Pressure Test' | pressure: 101325.0 | Researcher: Dr. Smith
[2026-04-03] ID=3 | 'Semiconductor Threshold' | voltage: 3.32 | Researcher: Dr. Frankov
[2026-04-04] ID=4 | 'Laser Cooling Chamber' | temperature: -210.15 | Researcher: Dr. Frankov

=== Temperature SQL Aggregates (SQLAlchemy) ===
Count: 2.0 | Average: -92.83 | Min: -210.15 | Max: 24.50

=== Raw DB-API Parameterized Search (sqlite3) ===
Found 3 experiments by Dr. Frankov:
  - [2026-04-01] Cryogenic Calibration (temperature = 24.5) | Notes: Initial calibration at room temperature
  - [2026-04-03] Semiconductor Threshold (voltage = 3.32) | Notes: Gate voltage stability
  - [2026-04-04] Laser Cooling Chamber (temperature = -210.15) | Notes: Deep cryocooling phase

DB-API Range Filter [-250.0, 50.0] C: 2 records found.
Global DB-API Stats: Count=4.0 | Avg=25285.67

=== Testing ACID Transaction Rollback ===
Caught expected transaction error: 'Physical constraint violation: Temperature -300.0 < -273.15 C'
Session successfully rolled back. Verifying count...
Initial count: 4.0 | Post-rollback count: 4.0 (No partial commits!)

All database demonstration steps completed successfully."""
    add_code_block(run_log)

    add_h2("6.2. Протокол виконання тестів pytest та Coverage Report")
    add_p("Усі 15 розроблених інтеграційних тестів виконано успішно за 0.78 с із досягненням 91 % загального покриття:")
    cov_log = """============================= test session starts =============================
platform win32 -- Python 3.13.12, pytest-9.1.1, pluggy-1.7.0
rootdir: Lab_07, configfile: pyproject.toml
plugins: cov-7.1.0
collected 15 items

Lab_07/tests/test_db_api.py .....                                        [ 33%]
Lab_07/tests/test_main.py .                                              [ 40%]
Lab_07/tests/test_migrations.py .                                        [ 46%]
Lab_07/tests/test_repositories.py .....                                  [ 80%]
Lab_07/tests/test_services.py ...                                        [100%]

=============================== tests coverage ================================
Name                                  Stmts   Miss Branch BrPart  Cover
-----------------------------------------------------------------------
experiment_manager/__init__.py            1      0      0      0   100%
experiment_manager/database.py           34      3      2      0    92%
experiment_manager/db_api.py             49      2     12      3    92%
experiment_manager/main.py               63      2      6      1    96%
experiment_manager/models.py             28      4      0      0    86%
experiment_manager/repositories.py       83      5     20      9    86%
experiment_manager/services.py           39      2      8      1    94%
-----------------------------------------------------------------------
TOTAL                                   297     18     48     14    91%
======================== 15 passed in 0.78s ========================"""
    add_code_block(cov_log)

    add_h2("6.3. Історія комітів Git")
    git_log_str = """fc1dbe0 Add Alembic database migrations and comprehensive integration test suite
f1806dd Add transactional business services and CLI application entrypoint
10783d5 Implement repository persistence layer and DB-API parameterized queries
39b154a Add SQLAlchemy database engine, session, and ORM domain models
c2810d1 Setup project structure, configuration, and alembic setup for Lab 07"""
    add_code_block(git_log_str)

    # ==================== ВІДПОВІДІ НА КОНТРОЛЬНІ ЗАПИТАННЯ ====================
    add_h1("7. ВІДПОВІДІ НА ВСІ КОНТРОЛЬНІ ЗАПИТАННЯ")

    q_and_a = [
        ("Що таке база даних?", "База даних — це структурована, організована сукупність даних, що зберігається в електронному вигляді та керується спеціальним програмним забезпеченням для швидкого пошуку, додавання й аналізу."),
        ("Що таке реляційна база даних?", "Це тип бази даних, що базується на реляційній моделі даних Е. Кодда, де інформація представлена у вигляді двовимірних взаємопов'язаних таблиць."),
        ("Що таке таблиця?", "Таблиця (відношення) — це основний елемент реляційної бази даних, що складається зі стовпчиків фіксованих типів та рядків записів."),
        ("Що таке рядок таблиці?", "Рядок (кортеж або запис) — це одиничний екземпляр даних у таблиці, що містить повний набір значень усіх полів для одного конкретного об'єкта."),
        ("Що таке поле?", "Поле (атрибут або стовпець) — це іменована складова таблиці з визначеним типом даних, що описує конкретну властивість сутності."),
        ("Що таке SQL?", "SQL (Structured Query Language) — стандартизована декларативна мова програмування для створення, модифікації та керування даними в реляційних СУБД."),
        ("Для чого використовується CREATE TABLE?", "Команда DDL CREATE TABLE використовується для визначення нової таблиці в базі даних, оголошення її стовпців, їх типів та обмежень цілісності."),
        ("Що означає CRUD?", "Акронім CRUD описує чотири фундаментальні операції роботи зі сховищем даних: Create (створення), Read (читання), Update (оновлення), Delete (видалення)."),
        ("Яка SQL-команда відповідає Create?", "Команда INSERT INTO."),
        ("Яка команда відповідає Read?", "Команда SELECT."),
        ("Яка команда відповідає Update?", "Команда UPDATE ... SET."),
        ("Яка команда відповідає Delete?", "Команда DELETE FROM."),
        ("Що таке SQLite?", "SQLite — це компактна, вбудована (serverless), безконфігураційна реляційна СУБД, яка зберігає всю базу даних в одному звичайному дисковому файлі."),
        ("Які переваги SQLite?", "Не потребує окремого сервера чи адміністрування, надзвичайно висока швидкість роботи, нульова конфігурація, підтримка повного ACID та наявність у стандартній бібліотеці Python (модуль sqlite3)."),
        ("Що таке Python DB-API?", "Python DB-API (PEP 249) — стандартизований інтерфейс для взаємодії програм на Python із різними реляційними СУБД через об'єкти Connection та Cursor."),
        ("Для чого використовується Connection?", "Об'єкт Connection представляє активне з'єднання з базою даних, керує ізоляцією транзакцій та надає методи commit() і rollback()."),
        ("Для чого використовується Cursor?", "Cursor — об'єкт, що дозволяє виконувати SQL-запити (метод execute()), обробляти результати вибірки (fetchall(), fetchone()) та відстежувати стан ітератора."),
        ("Для чого потрібен commit()?", "Метод commit() фіксує всі поточні зміни поточної транзакції в базі даних, роблячи їх постійними та видимими іншим з'єднанням."),
        ("Для чого потрібен rollback()?", "Метод rollback() скасовує всі зміни, внесені в межах поточної транзакції з моменту останнього commit(), повертаючи БД у початковий узгоджений стан."),
        ("Що таке SQL injection?", "SQL Injection (ін'єкція SQL) — небезпечна вразливість, коли зловмисник впроваджує довільні SQL-команди у вхідні дані застосунку, перехоплюючи або змінюючи логіку виконання запиту."),
        ("Що таке parameterized query?", "Параметризований запит — це техніка виконання SQL, де значення відокремлені від тексту запиту плейсхолдерами (?, :name) і передаються драйверу СУБД окремо як параметри."),
        ("Чому не рекомендується формувати SQL через f-string?", "Формування SQL через f-strings відкриває пряму критичну вразливість до SQL-ін'єкцій та спричиняє синтаксичні помилки при наявності лапок чи спецсимволів."),
        ("Що таке Primary Key?", "Primary Key (первинний ключ) — стовпець або набір стовпців, значення яких однозначно ідентифікують кожен рядок таблиці та не можуть містити NULL."),
        ("Що таке Foreign Key?", "Foreign Key (зовнішній ключ) — стовпець, значення якого посилається на Primary Key іншої таблиці, забезпечуючи реляційну цілісність зв'язків між сутностями."),
        ("Що таке транзакція?", "Транзакція — це неподільна послідовність операцій над базою даних, яка розглядається системою як єдине ціле і виконується за принципами ACID."),
        ("Що означає Atomicity?", "Атомарність гарантує правило «все або нічого»: або всі операції транзакції успішно завершуються, або при збої всі вони скасовуються (rollback)."),
        ("Що означає Consistency?", "Узгодженість гарантує, що транзакція переводить базу даних виключно з одного валідного стану в інший, не порушуючи правил, типів і зовнішніх ключів."),
        ("Що означає Isolation?", "Ізольованість гарантує, що проміжні стани паралельних транзакцій не впливають одна на одну до моменту їх фіксації."),
        ("Що означає Durability?", "Довговічність гарантує, що після успішного підтвердження транзакції (commit) зміни надійно збережені на диску навіть при раптовому вимкненні живлення."),
        ("Що означає ACID?", "Акронім фундаментальних властивостей надійних транзакцій у СУБД: Atomicity, Consistency, Isolation, Durability."),
        ("Що таке persistence layer?", "Persistence layer (шар збереження даних) — архітектурний рівень додатку, відповідальний за переміщення, збереження та отримання доменних об'єктів із довготривалого сховища."),
        ("Для чого потрібний Repository Pattern?", "Він абстрагує бізнес-логіку від деталей доступу до даних, емулюючи колекцію об'єктів у пам'яті та усуваючи пряме змішування SQL із доменними правилами."),
        ("Чому Repository відокремлюють від Service Layer?", "Щоб забезпечити Single Responsibility Principle: сервіси відповідають за бізнес-процеси, а репозиторії — виключно за операції збереження та пошуку даних."),
        ("Що таке ORM?", "Object-Relational Mapping (об'єктно-реляційне відображення) — технологія, що автоматично зіставляє таблиці реляційної БД з класами мови програмування."),
        ("Для чого використовується SQLAlchemy?", "SQLAlchemy — провідний ORM-інструмент та SQL Toolkit для Python, що надає високорівневе об'єктне моделювання, Unit of Work та генерацію оптимізованих запитів."),
        ("Що таке ORM model?", "ORM model — це Python-клас, успадкований від Base, атрибути якого відповідають стовпчикам певної реляційної таблиці."),
        ("Що таке relationship()?", "Функція SQLAlchemy, що зв'язує дві ORM-моделі через зовнішній ключ, дозволяючи звертатися до зв'язаних сутностей як до звичайних Python-атрибутів чи списків."),
        ("Що означає one-to-many?", "Відношення «один до багатьох», коли один рядок батьківської таблиці може бути пов'язаний із багатьма рядками дочірньої таблиці (наприклад, один дослідник і багато експериментів)."),
        ("Що означає many-to-many?", "Відношення «багато до багатьох», коли багато рядків однієї таблиці пов'язані з багатьма рядками іншої (зазвичай реалізується через проміжну асоціативну таблицю)."),
        ("Що таке migration?", "Міграція — це файл коду з версіонованими інструкціями оновлення структури схеми бази даних (додавання полів, зміна типів) без втрати існуючих даних."),
        ("Для чого використовується Alembic?", "Alembic — це інструмент керування міграціями бази даних для SQLAlchemy, що відстежує ревізії схеми та автоматизує процеси upgrade й downgrade."),
        ("Для чого потрібна команда alembic upgrade head?", "Команда alembic upgrade head застосовує всі новіші не застосовані міграції по черзі, оновлюючи схему БД до найостаннішої актуальної ревізії."),
        ("Що таке integration test?", "Інтеграційний тест — тест, що верифікує спільну роботу кількох компонентів додатку разом із реальними зовнішніми залежностями (базою даних або файлами)."),
        ("Чим integration test відрізняється від unit test?", "Unit-тест ізолює функцію за допомогою моків, тоді як інтеграційний тест перевіряє реальну взаємодію сервісів із реальною СУБД."),
        ("Навіщо використовувати окрему test database?", "Щоб тести не видаляли та не спотворювали робочі дані розробника або продуктового середовища, а також для гарантії чистого початкового стану перед кожним тестом."),
        ("Що перевіряє тест rollback?", "Тест rollback перевіряє виконання принципу атомарності (Atomicity): у разі виникнення помилки під час транзакції жоден рядок не залишається збереженим."),
        ("Для чого використовуються fixtures?", "Для стандартизованої підготовки тестового оточення (створення тестової сесії, генерація даних) та автоматичного прибирання за собою (teardown)."),
        ("Що робить SQL JOIN?", "Оператор JOIN об'єднує стовпчики з двох або більше таблиць в один результуючий набір даних на основі спільного логічного зв'язку (наприклад, researcher_id = researchers.id)."),
        ("Що робить GROUP BY?", "Оператор GROUP BY групує рядки, що мають однакові значення у визначених стовпцях, дозволяючи виконувати обчислення агрегатних функцій для кожної групи окремо."),
        ("Для чого використовуються COUNT, SUM, AVG, MIN, MAX?", "Це стандартні агрегатні функції SQL: COUNT підраховує кількість рядків, SUM обчислює суму, AVG — середнє арифметичне, MIN та MAX знаходять відповідно найменше та найбільше значення у стовпці.")
    ]

    for idx, (q, a) in enumerate(q_and_a, 1):
        add_p(f"{idx}. {q}", bold=True, space_after=3)
        add_p(f"Відповідь: {a}", space_after=8)

    # ==================== ВИСНОВОК ====================
    add_h1("8. ВИСНОВОК")
    add_p("У ході виконання лабораторної роботи № 7 було всебічно опановано сучасні стандарти та практики побудови надійного рівня збереження даних (Persistence Layer) у мові програмування Python.")
    add_p("Основні практичні результати роботи:")
    add_p("1. Спроєктовано та реалізовано реляційну схему бази даних SQLite для індивідуального Варіанта № 15 («Облік результатів наукових експериментів») із таблицями researchers та experiments, зв'язком One-to-Many та каскадним видаленням.")
    add_p("2. Застосовано патерн Repository (ResearcherRepository та ExperimentRepository), що забезпечило ізоляцію бізнес-логіки від особливостей СУБД та реалізувало операції CRUD, фільтрацію за діапазоном і сортування за датою.")
    add_p("3. Реалізовано безпечні параметризовані SQL-запити на базі стандарту Python DB-API (модуль sqlite3), що унеможливило вразливість до SQL Injection та забезпечило швидкий розрахунок агрегатних показників AVG, MIN, MAX, COUNT.")
    add_p("4. Реалізовано транзакційний сервісний шар із підтримкою пакетного додавання експериментів та верифікацією принципу ACID Atomicity за допомогою гарантованого відкату rollback() у разі виявлення фізично невалідних значень.")
    add_p("5. Налаштовано систему міграцій бази даних Alembic та створено скрипт 001_add_notes_to_experiments.py для безпечної еволюції схеми в режимі batch_alter_table.")
    add_p("6. Розроблено набір із 15 автоматизованих інтеграційних тестів на базі pytest із використанням тестової SQLite БД в оперативній пам'яті, що підтвердили 100 % функціональну коректність системи та забезпечили 91 % тестового покриття.")

    # Save document
    report_dir = Path("Lab_07/report")
    report_dir.mkdir(parents=True, exist_ok=True)
    report_file = report_dir / "ПП_ЛАБ7_ФРАНКОВ.docx"
    doc.save(str(report_file))
    print(f"Report successfully saved to: {report_file}")


if __name__ == "__main__":
    create_report()
